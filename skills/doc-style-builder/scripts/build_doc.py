#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx>=1.1", "playwright>=1.40"]
# ///
"""Markdown -> house-style PDF + DOCX (and HTML on request), all styled from assets/theme.css.

    uv run scripts/build_doc.py doc.md                      # -> "doc - Update <ddmmyyyy>.pdf" + ".docx" beside doc.md
    uv run scripts/build_doc.py doc.md --prefix ACME --out dist --formats html,pdf,docx
    uv run scripts/build_doc.py doc.md --no-date            # -> "doc.pdf" + "doc.docx"

Needs pandoc on PATH. PDF needs Chromium for playwright once: `uv run --with playwright playwright install chromium`.

Table rows whose first cell starts with [+] / [~] / [-] are highlighted as added / edited / removed
(removed is also struck through); the marker itself is removed from the output.
"""
import argparse
import datetime as dt
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
CSS = ROOT / "assets" / "theme.css"
MARKS = {"+": "added", "~": "edited", "-": "removed"}
MARK_RE = re.compile(r"^\[([+~-])\]\s*")
A4_W, A4_H, MARGIN_X, MARGIN_Y = 210, 297, 20, 18  # mm, same as @page in theme.css


def tokens() -> dict[str, str]:
    return {k: v.strip() for k, v in re.findall(r"--([\w-]+):\s*([^;]+);", CSS.read_text(encoding="utf-8"))}


def first_font(stack: str) -> str:
    return stack.split(",")[0].strip().strip("'\"")


def pt(value: str) -> float:
    return float(value.removesuffix("pt"))


def hexcolor(value: str) -> str:
    return value.lstrip("#").upper()


def pandoc(*args: str) -> None:
    # --columns=1: every table takes its column widths from the dash counts of its separator line
    # no syntax colors: the theme keeps color for meaning (pass/fail, change state)
    subprocess.run(["pandoc", "--columns=1", "--syntax-highlighting=none", *args], check=True)


# ---------- HTML + PDF ----------

def mark_html_rows(html: str) -> str:
    """<tr class="odd"><td>[+] x  ->  <tr class="odd added"><td>x"""
    pat = re.compile(r'<tr( class="([^"]*)")?>(\s*<td[^>]*>(?:\s*<p>)?)\[([+~-])\]\s*')
    return pat.sub(lambda m: f'<tr class="{((m[2] or "") + " " + MARKS[m[4]]).strip()}">{m[3]}', html)


def build_html(src: Path, out: Path) -> None:
    pandoc(str(src), "-f", "markdown", "-t", "html5", "-s", "--embed-resources", "--css", str(CSS),
           "--resource-path", str(src.parent), "-M", "lang=th", "-o", str(out))
    out.write_text(mark_html_rows(out.read_text(encoding="utf-8")), encoding="utf-8")


def build_pdf(html: Path, out: Path) -> None:
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        sys.exit("PDF needs playwright: run with `uv run`, or pass --formats docx")
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.goto(html.as_uri())
        page.evaluate("document.fonts.ready")
        page.pdf(path=str(out), format="A4", print_background=True, prefer_css_page_size=True)
        browser.close()


# ---------- DOCX ----------

def set_fonts(style, face: str) -> None:
    from docx.oxml.ns import qn
    style.font.name = face
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, fonts)
    for attr in list(fonts.attrib):  # theme fonts (asciiTheme, cstheme …) beat explicit ones, so drop them
        if attr.lower().endswith("theme"):
            del fonts.attrib[attr]
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):  # Thai is the complex script (cs)
        fonts.set(qn(k), face)


def set_bold(rpr) -> None:
    """bold for Latin (b) and for Thai, which Word draws as complex script (bCs)"""
    rpr.get_or_add_b()  # python-docx inserts both in ECMA-376 order
    rpr.get_or_add_bCs()


def set_size(style, size: float) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    cs = rpr.find(qn("w:szCs"))
    if cs is None:  # szCs follows sz in ECMA-376 order
        cs = rpr.makeelement(qn("w:szCs"), {})
        rpr.find(qn("w:sz")).addnext(cs)
    cs.set(qn("w:val"), str(int(size * 2)))


def reference_docx(path: Path, t: dict[str, str]) -> Path:
    """pandoc's default reference.docx restyled from the theme tokens"""
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import RGBColor
    with open(path, "wb") as fh:
        subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"], stdout=fh, check=True)
    doc = Document(str(path))
    if "Source Code" not in [st.name for st in doc.styles]:  # pandoc cites it but only defines it when highlighting
        doc.styles.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH).base_style = doc.styles["Normal"]
    body, head, code = first_font(t["font-body"]), first_font(t["font-heading"]), first_font(t["font-code"])
    sizes = {"Normal": "size-body", "Body Text": "size-body", "First Paragraph": "size-body", "Compact": "size-body",
             "Title": "size-title", "Subtitle": "size-subtitle", "Author": "size-body", "Date": "size-body",
             "Heading 1": "size-h1", "Heading 2": "size-h2", "Heading 3": "size-h3",
             "Image Caption": "size-caption", "Table Caption": "size-caption",
             "Source Code": "size-code-block", "Verbatim Char": "size-code"}
    for st in doc.styles:
        if getattr(st, "font", None) is None:
            continue
        is_code = st.name in ("Source Code", "Verbatim Char")
        is_head = st.name.startswith("Heading") or st.name == "Title"
        set_fonts(st, code if is_code else head if is_head else body)
        if st.name in sizes:
            set_size(st, pt(t[sizes[st.name]]))
        if is_head or st.name == "Subtitle":
            set_bold(st.element.get_or_add_rPr())
            st.font.color.rgb = RGBColor.from_string(hexcolor(t["color-text"]))
        if st.name in ("Title", "Subtitle", "Author", "Date", "Image Caption"):
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if st.name == "Image Caption":
            st.font.italic = True
            st.font.color.rgb = RGBColor.from_string(hexcolor(t["color-caption"]))
    box_code_style(doc.styles["Source Code"], t)
    doc.save(str(path))
    return path


def shade(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(qn("w:shd")):
        tcpr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def strip_mark(cell) -> str | None:
    """remove a leading [+]/[~]/[-] from the cell; return the change name it carried"""
    for para in cell.paragraphs:
        runs = [r for r in para.runs if r.text]
        if not runs:
            continue
        text = "".join(r.text for r in runs)
        m = MARK_RE.match(text)
        if not m:
            return None
        cut = m.end()
        for r in runs:  # the marker may span several runs
            take = min(cut, len(r.text))
            r.text, cut = r.text[take:], cut - take
            if cut == 0:
                break
        return MARKS[m[1]]
    return None


def style_tables(doc, t: dict[str, str]) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    text_twips = int((A4_W - 2 * MARGIN_X) / 25.4 * 1440)
    fills = {name: hexcolor(t[f"color-{name}"]) for name in MARKS.values()}
    for table in doc.tables:
        tblpr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", hexcolor(t["color-rule"]))):
                e.set(qn(f"w:{k}"), v)
            borders.append(e)
        tblpr.append(borders)
        for old in tblpr.findall(qn("w:tblW")):
            tblpr.remove(old)
        width = OxmlElement("w:tblW"); width.set(qn("w:w"), "5000"); width.set(qn("w:type"), "pct")
        tblpr.append(width)
        cols = table._tbl.tblGrid.findall(qn("w:gridCol"))  # pandoc's fixed grid wins over pct: scale it to text width
        total = sum(int(c.get(qn("w:w"))) for c in cols) or 1
        for c in cols:
            c.set(qn("w:w"), str(int(int(c.get(qn("w:w"))) * text_twips / total)))
        for tcw in list(table._tbl.iter(qn("w:tcW"))):
            tcw.getparent().remove(tcw)
        has_header = bool(table._tbl.xpath(".//w:trPr/w:tblHeader"))  # pandoc marks header rows; an empty header = key/value table
        for row in table.rows:
            trpr = row._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))  # a row never splits across pages
            if trpr.find(qn("w:tblHeader")) is not None:
                for cell in row.cells:
                    shade(cell, hexcolor(t["color-header-bg"]))
                    for para in cell.paragraphs:
                        for r in para.runs:
                            set_bold(r._r.get_or_add_rPr())
                continue
            if not has_header:
                shade(row.cells[0], hexcolor(t["color-header-bg"]))
            change = strip_mark(row.cells[0])
            if change:
                for cell in row.cells:
                    shade(cell, fills[change])
                    if change == "removed":
                        for para in cell.paragraphs:
                            for r in para.runs:
                                r.font.strike = True


def box_code_style(style, t: dict[str, str]) -> None:
    """code blocks: bordered and shaded like the <pre> in theme.css"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    ppr = style.element.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        for k, v in (("val", "single"), ("sz", "4"), ("space", "4"), ("color", hexcolor(t["color-border"]))):
            e.set(qn(f"w:{k}"), v)
        border.append(e)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor(t["color-code-bg"]))
    # ECMA-376 orders pPr children; pBdr then shd come after these, so insert past the last one present
    before = {qn(f"w:{n}") for n in ("pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
                                       "widowControl", "numPr", "suppressLineNumbers")}
    at = max((i + 1 for i, c in enumerate(ppr) if c.tag in before), default=0)
    ppr.insert(at, shd)
    ppr.insert(at, border)


def build_docx(src: Path, out: Path, work: Path, t: dict[str, str]) -> None:
    from docx import Document
    from docx.shared import Mm, Pt
    ref = reference_docx(work / "reference.docx", t)
    pandoc(str(src), "-f", "markdown", "-t", "docx", "--reference-doc", str(ref),
           "--resource-path", str(src.parent), "-M", "lang=th", "-o", str(out))
    doc = Document(str(out))
    style_tables(doc, t)
    code_size = Pt(pt(t["size-code-block"]))
    for para in doc.paragraphs:  # code runs carry Verbatim Char (inline size); blocks are smaller
        if para.style.name == "Source Code":
            for run in para.runs:
                run.font.size = code_size
    for sec in doc.sections:
        sec.page_width, sec.page_height = Mm(A4_W), Mm(A4_H)
        sec.left_margin = sec.right_margin = Mm(MARGIN_X)
        sec.top_margin = sec.bottom_margin = Mm(MARGIN_Y)
    max_w = Mm(A4_W - 2 * MARGIN_X)
    for shape in doc.inline_shapes:
        if shape.width > max_w:
            shape.height, shape.width = int(shape.height * max_w / shape.width), max_w
    doc.save(str(out))


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("src", type=Path, help="Markdown source (front matter: title, subtitle, author, date)")
    ap.add_argument("--out", type=Path, help="output folder (default: beside the source)")
    ap.add_argument("--prefix", default="", help="file-name prefix, e.g. a client or system code")
    ap.add_argument("--no-date", action="store_true", help='omit " - Update <ddmmyyyy>" from the file name')
    ap.add_argument("--formats", default="pdf,docx", help="comma list of html, pdf, docx")
    args = ap.parse_args()

    if not shutil.which("pandoc"):
        sys.exit("pandoc not found on PATH")
    src = args.src.resolve()
    formats = {f.strip() for f in args.formats.split(",") if f.strip()}
    if not formats <= {"html", "pdf", "docx"}:
        sys.exit(f"unknown format: {', '.join(formats - {'html', 'pdf', 'docx'})}")
    out_dir = (args.out or src.parent).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    name = (f"{args.prefix}_" if args.prefix else "") + src.stem
    if not args.no_date:
        name += f" - Update {dt.date.today():%d%m%Y}"
    t = tokens()

    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        if formats & {"html", "pdf"}:
            html = (out_dir if "html" in formats else work) / f"{name}.html"
            build_html(src, html)
            if "html" in formats:
                print(html)
            if "pdf" in formats:
                pdf = out_dir / f"{name}.pdf"
                build_pdf(html, pdf)
                print(pdf)
        if "docx" in formats:
            docx = out_dir / f"{name}.docx"
            build_docx(src, docx, work, t)
            print(docx)


if __name__ == "__main__":
    main()
